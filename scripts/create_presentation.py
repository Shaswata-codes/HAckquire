import sys
import os

def create_deck():
    from pptx import Presentation
    from pptx.util import Inches, Pt
    from pptx.dml.color import RGBColor
    from pptx.enum.text import PP_ALIGN
    from pptx.enum.shapes import MSO_SHAPE

    prs = Presentation()
    prs.slide_width = Inches(13.333)  # 16:9 widescreen
    prs.slide_height = Inches(7.5)

    # Color Palette
    BG_COLOR = RGBColor(11, 15, 25)       # Slate 950 #0b0f19
    CARD_BG = RGBColor(22, 30, 49)        # Slate 900 #161e31
    BORDER_COLOR = RGBColor(40, 53, 84)   # Slate 800
    TEXT_WHITE = RGBColor(248, 250, 252)  # White
    TEXT_MUTED = RGBColor(148, 163, 184)  # Slate 400
    INDIGO_ACCENT = RGBColor(99, 102, 241)# Indigo #6366f1
    CYAN_ACCENT = RGBColor(6, 182, 212)   # Cyan #06b6d4
    EMERALD_ACCENT = RGBColor(16, 185, 129)# Emerald #10b981
    GOLD_ACCENT = RGBColor(245, 158, 11)  # Amber

    def set_slide_bg(slide):
        background = slide.background
        fill = background.fill
        fill.solid()
        fill.fore_color.rgb = BG_COLOR

    blank_layout = prs.slide_layouts[6]

    # =========================================================================
    # SLIDE 1: TEAM & PROJECT OVERVIEW
    # =========================================================================
    slide1 = prs.slides.add_slide(blank_layout)
    set_slide_bg(slide1)

    # Top Tag
    tb = slide1.shapes.add_textbox(Inches(0.8), Inches(0.6), Inches(11.7), Inches(0.5))
    tf = tb.text_frame
    p = tf.paragraphs[0]
    p.text = "HACKATHON 2026 • FINTECH & AI AUTOMATION"
    p.font.size = Pt(11)
    p.font.bold = True
    p.font.color.rgb = CYAN_ACCENT

    # Main Title
    tb_title = slide1.shapes.add_textbox(Inches(0.8), Inches(1.0), Inches(11.7), Inches(1.1))
    tf_title = tb_title.text_frame
    p_title = tf_title.paragraphs[0]
    p_title.text = "HACKQUIRE"
    p_title.font.size = Pt(44)
    p_title.font.bold = True
    p_title.font.color.rgb = TEXT_WHITE

    # Subtitle
    tb_sub = slide1.shapes.add_textbox(Inches(0.8), Inches(2.0), Inches(11.7), Inches(0.7))
    tf_sub = tb_sub.text_frame
    p_sub = tf_sub.paragraphs[0]
    p_sub.text = "AI-Powered Vernacular Invoicing & Automated Financial Operations for the Gig Economy"
    p_sub.font.size = Pt(16)
    p_sub.font.color.rgb = INDIGO_ACCENT

    # Value Prop Box
    box_vp = slide1.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(0.8), Inches(2.8), Inches(11.7), Inches(1.1))
    box_vp.fill.solid()
    box_vp.fill.fore_color.rgb = CARD_BG
    box_vp.line.color.rgb = BORDER_COLOR
    tf_vp = box_vp.text_frame
    tf_vp.word_wrap = True
    p_vp = tf_vp.paragraphs[0]
    p_vp.text = "🎯 Core Value Proposition:"
    p_vp.font.size = Pt(12)
    p_vp.font.bold = True
    p_vp.font.color.rgb = EMERALD_ACCENT
    p_vp2 = tf_vp.add_paragraph()
    p_vp2.text = "Empowers gig workers, tutors, tailors, electricians, and freelancers to generate GST-ready professional invoices in 3 seconds by speaking in regional Indian languages (Hindi, Bengali, Tamil, etc.) and collect payments 3x faster with 1-click UPI reconciliation."
    p_vp2.font.size = Pt(11)
    p_vp2.font.color.rgb = TEXT_MUTED

    # Team Members (4 Columns)
    team_members = [
        {"name": "Shaswata Sarkar", "id": "24155125", "role": "Team Lead & Full-Stack Architect", "color": INDIGO_ACCENT},
        {"name": "Soham Chakraborty", "id": "24155502", "role": "Backend & AI Systems Engineer", "color": CYAN_ACCENT},
        {"name": "Tanima Maity", "id": "2455138", "role": "Frontend & UI/UX Specialist", "color": EMERALD_ACCENT},
        {"name": "Soumyadip Jana", "id": "24155130", "role": "Integration & Cloud Deployments", "color": GOLD_ACCENT}
    ]

    col_w = Inches(2.75)
    col_gap = Inches(0.24)
    start_x = Inches(0.8)

    for i, m in enumerate(team_members):
        x = start_x + i * (col_w + col_gap)
        card = slide1.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, x, Inches(4.3), col_w, Inches(2.5))
        card.fill.solid()
        card.fill.fore_color.rgb = CARD_BG
        card.line.color.rgb = BORDER_COLOR
        
        tf_c = card.text_frame
        tf_c.word_wrap = True
        
        p1 = tf_c.paragraphs[0]
        p1.text = m["name"]
        p1.font.size = Pt(14)
        p1.font.bold = True
        p1.font.color.rgb = TEXT_WHITE
        
        p2 = tf_c.add_paragraph()
        p2.text = f"ID: {m['id']}"
        p2.font.size = Pt(11)
        p2.font.bold = True
        p2.font.color.rgb = m["color"]
        
        p3 = tf_c.add_paragraph()
        p3.text = f"\n{m['role']}"
        p3.font.size = Pt(10)
        p3.font.color.rgb = TEXT_MUTED

    # =========================================================================
    # SLIDE 2: PROTOTYPE ARCHITECTURE & AI ENGINE
    # =========================================================================
    slide2 = prs.slides.add_slide(blank_layout)
    set_slide_bg(slide2)

    # Slide Header
    tb2 = slide2.shapes.add_textbox(Inches(0.8), Inches(0.5), Inches(11.7), Inches(0.9))
    tf2 = tb2.text_frame
    p = tf2.paragraphs[0]
    p.text = "PROTOTYPE OVERVIEW (PART 1)"
    p.font.size = Pt(10)
    p.font.bold = True
    p.font.color.rgb = CYAN_ACCENT
    p2 = tf2.add_paragraph()
    p2.text = "Multilingual AI Invoicing & Indic Intelligence Engine"
    p2.font.size = Pt(22)
    p2.font.bold = True
    p2.font.color.rgb = TEXT_WHITE

    # 3 Core Architecture Pillars
    pillars = [
        {
            "icon": "🎙️ AI Multilingual Invoice Generator",
            "tag": "Gemini 1.5 Flash + Indic Heuristics",
            "tag_color": INDIGO_ACCENT,
            "points": [
                "Understands unstructured natural language in 10+ regional Indian languages (Hindi, Bengali, Tamil, Telugu, Marathi, Hinglish).",
                "Extracts Client Name, Line Items, Quantities, Unit Rates, and Due Dates automatically.",
                "Real Example: 'প্রিয়ার অঙ্ক টিউশন ১০ টি ক্লাস ৫০০ টাকা' ➔ Extracted: Client Priya | 10 Classes @ ₹500/class | Total ₹5,000.",
                "Zero-failure architecture with offline heuristic fallback parser."
            ]
        },
        {
            "icon": "📄 High-Fidelity Client-Side PDF Engine",
            "tag": "100% Native Indic Script Support",
            "tag_color": EMERALD_ACCENT,
            "points": [
                "2x Retina-scale client-side rendering with zero server render latency.",
                "Full UTF-8 Unicode support for Bengali, Hindi, Tamil, Telugu scripts and ₹ currency symbol.",
                "Instant 1-click professional branded PDF download with tax, discounts, and breakdown."
            ]
        },
        {
            "icon": "🔐 Identity & Security Sandbox",
            "tag": "Google OAuth 2.0 + Instant Demo",
            "tag_color": CYAN_ACCENT,
            "points": [
                "Native Google OAuth 2.0 popup with single-click account chooser on device.",
                "Pre-seeded 1-click sandbox demo (demo@hackquire.com / demo123) for instant hackathon evaluation.",
                "Encrypted JWT authentication and secured RESTful endpoints."
            ]
        }
    ]

    card_w = Inches(3.72)
    card_gap = Inches(0.26)
    start_x = Inches(0.8)

    for i, p_data in enumerate(pillars):
        x = start_x + i * (card_w + card_gap)
        c = slide2.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, x, Inches(1.6), card_w, Inches(5.3))
        c.fill.solid()
        c.fill.fore_color.rgb = CARD_BG
        c.line.color.rgb = BORDER_COLOR
        
        tfc = c.text_frame
        tfc.word_wrap = True
        
        p = tfc.paragraphs[0]
        p.text = p_data["icon"]
        p.font.size = Pt(13)
        p.font.bold = True
        p.font.color.rgb = TEXT_WHITE
        
        p_sub = tfc.add_paragraph()
        p_sub.text = p_data["tag"]
        p_sub.font.size = Pt(9)
        p_sub.font.bold = True
        p_sub.font.color.rgb = p_data["tag_color"]
        
        tfc.add_paragraph().text = "" # spacer
        
        for pt in p_data["points"]:
            p_pt = tfc.add_paragraph()
            p_pt.text = f"• {pt}"
            p_pt.font.size = Pt(10)
            p_pt.font.color.rgb = TEXT_MUTED
            p_pt.space_after = Pt(6)

    # =========================================================================
    # SLIDE 3: FINANCIAL AUTOMATION & DUE ENGINE
    # =========================================================================
    slide3 = prs.slides.add_slide(blank_layout)
    set_slide_bg(slide3)

    tb3 = slide3.shapes.add_textbox(Inches(0.8), Inches(0.5), Inches(11.7), Inches(0.9))
    tf3 = tb3.text_frame
    p = tf3.paragraphs[0]
    p.text = "PROTOTYPE OVERVIEW (PART 2)"
    p.font.size = Pt(10)
    p.font.bold = True
    p.font.color.rgb = CYAN_ACCENT
    p2 = tf3.add_paragraph()
    p2.text = "Payment Automation, UPI Reconciliation & Notification Engine"
    p2.font.size = Pt(22)
    p2.font.bold = True
    p2.font.color.rgb = TEXT_WHITE

    fin_modules = [
        {
            "icon": "⚡ 1-Click UPI Payment Reconciliation",
            "tag": "Instant Transaction Matching",
            "tag_color": EMERALD_ACCENT,
            "points": [
                "Matches incoming UPI transaction IDs (GPay, PhonePe, Paytm) directly to pending invoices.",
                "Auto-calculates remaining balances and auto-transitions status: Paid, Partial, or Overdue.",
                "Manual dispute review & unmatch tools for irregular amounts or mismatched payments."
            ]
        },
        {
            "icon": "📲 Smart Due Reminders & Nudges",
            "tag": "1-Click WhatsApp & SMS Dispatch",
            "tag_color": INDIGO_ACCENT,
            "points": [
                "Real-time tracking of upcoming and overdue payment deadlines.",
                "Generates polite, customized WhatsApp reminder messages with invoice number and amount due.",
                "Built-in SMS dispatcher for instant mobile notifications without awkward phone calls."
            ]
        },
        {
            "icon": "🌐 B2B Integrable Microservice",
            "tag": "Exportable API Engine",
            "tag_color": GOLD_ACCENT,
            "points": [
                "Exposed reusable endpoint POST /api/ai/invoice ready for plug-and-play marketplace integration.",
                "Multilingual translation endpoints POST /api/ai/translate for third-party platforms.",
                "Production deployed on Vercel Full-Stack Serverless with unified domain & zero CORS issues."
            ]
        }
    ]

    for i, m_data in enumerate(fin_modules):
        x = start_x + i * (card_w + card_gap)
        c = slide3.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, x, Inches(1.6), card_w, Inches(5.3))
        c.fill.solid()
        c.fill.fore_color.rgb = CARD_BG
        c.line.color.rgb = BORDER_COLOR
        
        tfc = c.text_frame
        tfc.word_wrap = True
        
        p = tfc.paragraphs[0]
        p.text = m_data["icon"]
        p.font.size = Pt(13)
        p.font.bold = True
        p.font.color.rgb = TEXT_WHITE
        
        p_sub = tfc.add_paragraph()
        p_sub.text = m_data["tag"]
        p_sub.font.size = Pt(9)
        p_sub.font.bold = True
        p_sub.font.color.rgb = m_data["tag_color"]
        
        tfc.add_paragraph().text = "" # spacer
        
        for pt in m_data["points"]:
            p_pt = tfc.add_paragraph()
            p_pt.text = f"• {pt}"
            p_pt.font.size = Pt(10)
            p_pt.font.color.rgb = TEXT_MUTED
            p_pt.space_after = Pt(6)

    # =========================================================================
    # SLIDE 4: TRADABLE ASSETS: BUY & SELL MARKETPLACE ECONOMY
    # =========================================================================
    slide4 = prs.slides.add_slide(blank_layout)
    set_slide_bg(slide4)

    tb4 = slide4.shapes.add_textbox(Inches(0.8), Inches(0.5), Inches(11.7), Inches(0.9))
    tf4 = tb4.text_frame
    p = tf4.paragraphs[0]
    p.text = "MARKETPLACE TRANSACTIONS & ECOSYSTEM VALUE"
    p.font.size = Pt(10)
    p.font.bold = True
    p.font.color.rgb = CYAN_ACCENT
    p2 = tf4.add_paragraph()
    p2.text = "Tradable Asset Portfolio: What We Sold & Bought"
    p2.font.size = Pt(22)
    p2.font.bold = True
    p2.font.color.rgb = TEXT_WHITE

    # Left Column: ASSETS SOLD (2 Cards)
    left_w = Inches(5.6)
    c_sold = slide4.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(0.8), Inches(1.6), left_w, Inches(5.3))
    c_sold.fill.solid()
    c_sold.fill.fore_color.rgb = CARD_BG
    c_sold.line.color.rgb = EMERALD_ACCENT
    tfc_s = c_sold.text_frame
    tfc_s.word_wrap = True

    p = tfc_s.paragraphs[0]
    p.text = "💰 ASSETS SOLD (Monetized Services)"
    p.font.size = Pt(14)
    p.font.bold = True
    p.font.color.rgb = EMERALD_ACCENT

    items_sold = [
        ("1. AI Invoice Generator API", "SOLD 2 TIMES", INDIGO_ACCENT, "Provided external teams with natural language invoice parsing microservice (POST /api/ai/invoice) to convert informal text into structured line items."),
        ("2. WhatsApp & SMS Due Reminder Engine", "SOLD 1 TIME", CYAN_ACCENT, "Exported our automated payment reminder dispatcher for 1-click notification workflows and client payment alerts.")
    ]

    for title, badge, b_color, desc in items_sold:
        tfc_s.add_paragraph().text = ""
        p_t = tfc_s.add_paragraph()
        p_t.text = f"✔ {title}"
        p_t.font.size = Pt(12)
        p_t.font.bold = True
        p_t.font.color.rgb = TEXT_WHITE

        p_b = tfc_s.add_paragraph()
        p_b.text = f"   [{badge}]"
        p_b.font.size = Pt(10)
        p_b.font.bold = True
        p_b.font.color.rgb = b_color

        p_d = tfc_s.add_paragraph()
        p_d.text = f"   {desc}"
        p_d.font.size = Pt(10)
        p_d.font.color.rgb = TEXT_MUTED

    # Right Column: ASSETS BOUGHT (4 Integrations)
    c_bought = slide4.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(6.8), Inches(1.6), left_w, Inches(5.3))
    c_bought.fill.solid()
    c_bought.fill.fore_color.rgb = CARD_BG
    c_bought.line.color.rgb = INDIGO_ACCENT
    tfc_b = c_bought.text_frame
    tfc_b.word_wrap = True

    p = tfc_b.paragraphs[0]
    p.text = "📥 ASSETS BOUGHT & INTEGRATED"
    p.font.size = Pt(14)
    p.font.bold = True
    p.font.color.rgb = INDIGO_ACCENT

    items_bought = [
        ("1. OCR + ReCorrect Document Engine (SAHAYAK AI)", "OCR & field extraction for automated ID and paper receipt scanning."),
        ("2. LedgerFlow Financial App (Next.js)", "High-performance financial ledger UI architecture and font system."),
        ("3. Vernacular / Regional-Language Wrapper (Module C)", "Multi-language transliteration & phonetic pronunciation layer for 10+ languages."),
        ("4. Tech-Tide Energy Monitoring Engine", "Slab-based billing calculations, anomaly detection, and alert simulator.")
    ]

    for title, desc in items_bought:
        tfc_b.add_paragraph().text = ""
        p_t = tfc_b.add_paragraph()
        p_t.text = f"★ {title}"
        p_t.font.size = Pt(11)
        p_t.font.bold = True
        p_t.font.color.rgb = TEXT_WHITE

        p_d = tfc_b.add_paragraph()
        p_d.text = f"   {desc}"
        p_d.font.size = Pt(9.5)
        p_d.font.color.rgb = TEXT_MUTED

    out_path = os.path.abspath("Hackquire_Presentation.pptx")
    prs.save(out_path)
    print(f"SUCCESS: Created presentation at {out_path}")

if __name__ == "__main__":
    create_deck()
